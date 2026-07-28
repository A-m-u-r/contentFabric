from __future__ import annotations

import re
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parent
MARKDOWN = ROOT / "Анализ_литературы_генеративные_AI_Shorts.md"
DOCX = ROOT / "Анализ_литературы_генеративные_AI_Shorts.docx"


def extract_numbered_entries(text: str, prefix: str) -> list[int]:
    pattern = rf"^#{{2,3}} \[{re.escape(prefix)}(\d+)\]"
    return [int(value) for value in re.findall(pattern, text, flags=re.MULTILINE)]


def main() -> None:
    failures: list[str] = []
    source = MARKDOWN.read_text(encoding="utf-8")
    document = Document(DOCX)

    expected = {"RU": 16, "EN": 19, "STD": 3}
    for prefix, total in expected.items():
        values = extract_numbered_entries(source, prefix)
        wanted = list(range(1, total + 1))
        if values != wanted:
            failures.append(f"{prefix} entries: expected {wanted}, got {values}")

    all_labels = {
        (prefix, int(number))
        for prefix, number in re.findall(r"\[(RU|EN|STD)(\d+)\]", source)
    }
    for prefix, number in sorted(all_labels):
        if number < 1 or number > expected[prefix]:
            failures.append(f"Unknown citation [{prefix}{number}]")

    markdown_links = re.findall(r"\[[^\]]+\]\((https?://[^)]+)\)", source)
    if len(markdown_links) < 38:
        failures.append(f"Too few source hyperlinks: {len(markdown_links)}")

    forbidden = ("[[TOC]]", "TODO", "TBD", "ЗАПОЛНИТЬ")
    document_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    for token in forbidden:
        if token in source or token in document_text:
            failures.append(f"Unresolved placeholder: {token}")

    heading_counts = {
        name: sum(1 for paragraph in document.paragraphs if paragraph.style.name == name)
        for name in ("Heading 1", "Heading 2", "Heading 3")
    }
    if heading_counts["Heading 1"] < 14:
        failures.append(f"Too few Heading 1 paragraphs: {heading_counts['Heading 1']}")
    if heading_counts["Heading 2"] < 60:
        failures.append(f"Too few Heading 2 paragraphs: {heading_counts['Heading 2']}")
    if len(document.tables) != 1:
        failures.append(f"Expected one evidence table, got {len(document.tables)}")

    title_style = document.styles["Title"]
    if title_style.paragraph_format.alignment != WD_ALIGN_PARAGRAPH.LEFT:
        failures.append("Title style is not left-aligned")

    section = document.sections[0]
    section_geometry = (
        round(section.page_width.inches, 2),
        round(section.page_height.inches, 2),
        round(section.left_margin.inches, 2),
        round(section.right_margin.inches, 2),
        round(section.top_margin.inches, 2),
        round(section.bottom_margin.inches, 2),
    )
    if section_geometry != (8.5, 11.0, 1.0, 1.0, 0.82, 0.78):
        failures.append(f"Unexpected section geometry: {section_geometry}")

    footer_xml = section.footer._element.xml
    if " PAGE " not in footer_xml:
        failures.append("Footer page-number field missing")

    with zipfile.ZipFile(DOCX) as archive:
        bad_member = archive.testzip()
        if bad_member:
            failures.append(f"Corrupt ZIP member: {bad_member}")
        relationship_xml = archive.read("word/_rels/document.xml.rels").decode("utf-8")
        external_hyperlinks = relationship_xml.count('TargetMode="External"')
        if external_hyperlinks < 38:
            failures.append(f"Too few DOCX external hyperlinks: {external_hyperlinks}")

    print(f"RU sources: {expected['RU']}")
    print(f"EN sources: {expected['EN']}")
    print(f"Standards: {expected['STD']}")
    print(f"Markdown hyperlinks: {len(markdown_links)}")
    print(f"DOCX hyperlinks: {external_hyperlinks}")
    print(f"Headings: {heading_counts}")
    print(f"Tables: {len(document.tables)}")
    print(f"Section geometry: {section_geometry}")
    print(f"DOCX size: {DOCX.stat().st_size} bytes")

    if failures:
        for failure in failures:
            print(f"FAIL: {failure}")
        raise SystemExit(1)
    print("VERIFICATION: PASS")


if __name__ == "__main__":
    main()
