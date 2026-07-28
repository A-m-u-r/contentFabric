from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "Анализ_литературы_генеративные_AI_Shorts.md"
OUTPUT = ROOT / "Анализ_литературы_генеративные_AI_Shorts.docx"

ACCENT = "1F4E79"
ACCENT_2 = "2F5597"
DARK = "17365D"
TEXT = "243447"
MUTED = "65758B"
LIGHT = "F4F6F9"
PALE_BLUE = "EAF2F8"
BORDER = "CBD7E6"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color: str = BORDER, size: str = "5") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_table_width_and_indent(table, width_dxa: int = 9360, indent_dxa: int = 120) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_paragraph_border(paragraph, *, side: str, color: str, size: str = "8", space: str = "4") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    edge = p_bdr.find(qn(f"w:{side}"))
    if edge is None:
        edge = OxmlElement(f"w:{side}")
        p_bdr.append(edge)
    edge.set(qn("w:val"), "single")
    edge.set(qn("w:sz"), size)
    edge.set(qn("w:space"), space)
    edge.set(qn("w:color"), color)


def shade_paragraph(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_run_font(run, name: str = "Calibri", size: float | None = None, color: str | None = None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:cs"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    lang = run._element.rPr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        run._element.rPr.append(lang)
    lang.set(qn("w:val"), "ru-RU")
    lang.set(qn("w:eastAsia"), "ru-RU")


def add_hyperlink(paragraph, text: str, url: str, *, bold: bool = False) -> None:
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), ACCENT_2)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Calibri")
    r_fonts.set(qn("w:hAnsi"), "Calibri")
    r_fonts.set(qn("w:eastAsia"), "Calibri")
    r_fonts.set(qn("w:cs"), "Calibri")
    if bold:
        r_pr.append(OxmlElement("w:b"))
    r_pr.append(r_fonts)
    r_pr.append(color)
    r_pr.append(underline)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE_PATTERN = re.compile(r"(\*\*.+?\*\*|\[[^\]]+\]\(https?://[^)]+\)|`[^`]+`)")


def add_inline(paragraph, text: str, *, default_bold: bool = False, default_color: str | None = None) -> None:
    pos = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos : match.start()])
            set_run_font(run, color=default_color)
            run.bold = default_bold
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, color=default_color)
            run.bold = True
        elif token.startswith("["):
            label, url = re.match(r"\[([^\]]+)\]\((https?://[^)]+)\)", token).groups()
            add_hyperlink(paragraph, label, url, bold=default_bold)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name="Consolas", size=9.5, color=DARK)
            run.font.highlight_color = None
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, color=default_color)
        run.bold = default_bold


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, size=9, color=MUTED)


def create_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333
    normal.paragraph_format.widow_control = True

    for style_name, size, color, before, after in (
        ("Title", 28, DARK, 0, 12),
        ("Subtitle", 14, MUTED, 0, 10),
        ("Heading 1", 16, ACCENT, 18, 10),
        ("Heading 2", 13, ACCENT_2, 12, 6),
        ("Heading 3", 12, DARK, 8, 4),
    ):
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = style_name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.widow_control = True
        if style_name == "Title":
            style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            style.font.bold = True
            style.paragraph_format.line_spacing = 1.05
        elif style_name == "Subtitle":
            style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            style.paragraph_format.line_spacing = 1.15
        else:
            style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            style.paragraph_format.line_spacing = 1.08

    if "Report Meta" not in styles:
        meta = styles.add_style("Report Meta", WD_STYLE_TYPE.PARAGRAPH)
    else:
        meta = styles["Report Meta"]
    meta.font.name = "Calibri"
    meta._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    meta.font.size = Pt(10)
    meta.font.color.rgb = RGBColor.from_string(MUTED)
    meta.paragraph_format.space_after = Pt(4)
    meta.paragraph_format.line_spacing = 1.12

    if "Callout" not in styles:
        callout = styles.add_style("Callout", WD_STYLE_TYPE.PARAGRAPH)
    else:
        callout = styles["Callout"]
    callout.font.name = "Calibri"
    callout._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    callout.font.size = Pt(10.5)
    callout.font.color.rgb = RGBColor.from_string(DARK)
    callout.font.italic = True
    callout.paragraph_format.left_indent = Inches(0.22)
    callout.paragraph_format.right_indent = Inches(0.16)
    callout.paragraph_format.space_before = Pt(6)
    callout.paragraph_format.space_after = Pt(9)
    callout.paragraph_format.line_spacing = 1.22
    callout.paragraph_format.keep_together = True

    for list_name in ("List Bullet", "List Number"):
        style = styles[list_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(11)
        style.font.color.rgb = RGBColor.from_string(TEXT)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = -Inches(0.194)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.208
        style.paragraph_format.widow_control = True

    if "Table Text" not in styles:
        table_text = styles.add_style("Table Text", WD_STYLE_TYPE.PARAGRAPH)
    else:
        table_text = styles["Table Text"]
    table_text.font.name = "Calibri"
    table_text._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    table_text.font.size = Pt(8.5)
    table_text.font.color.rgb = RGBColor.from_string(TEXT)
    table_text.paragraph_format.space_after = Pt(2)
    table_text.paragraph_format.line_spacing = 1.08
    table_text.paragraph_format.widow_control = True


def configure_sections(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.78)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.38)
    section.different_first_page_header_footer = True

    first_header = section.first_page_header
    p = first_header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("ОБЗОР ЛИТЕРАТУРЫ  /  OPENSHORTS")
    set_run_font(r, size=8.5, color=ACCENT)
    r.bold = True
    set_paragraph_border(p, side="bottom", color=ACCENT, size="12", space="5")

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("МУЛЬТИМОДАЛЬНАЯ ГЕНЕРАЦИЯ AI SHORTS")
    set_run_font(r, size=8.5, color=MUTED)
    r.bold = True
    set_paragraph_border(p, side="bottom", color=BORDER, size="4", space="5")

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("OpenShorts  •  рабочий материал ВКР  •  ")
    set_run_font(r, size=9, color=MUTED)
    add_page_number(p)


def add_cover(doc: Document) -> None:
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(38)

    label = doc.add_paragraph()
    label.paragraph_format.space_after = Pt(12)
    run = label.add_run("АНАЛИТИЧЕСКИЙ ОТЧЁТ  /  ГЛАВА 1 ВКР")
    set_run_font(run, size=10, color=ACCENT)
    run.bold = True
    set_paragraph_border(label, side="bottom", color=ACCENT, size="18", space="8")

    title = doc.add_paragraph(style="Title")
    add_inline(title, "Анализ литературы по мультимодальной генерации коротких вертикальных видеороликов")

    subtitle = doc.add_paragraph(style="Subtitle")
    add_inline(subtitle, "Временная синхронизация сценария, речи, виртуального аватара, сцен и субтитров")

    statement = doc.add_paragraph(style="Callout")
    shade_paragraph(statement, PALE_BLUE)
    set_paragraph_border(statement, side="left", color=ACCENT, size="18", space="8")
    add_inline(
        statement,
        "Итог анализа: тема пригодна для ВКР при условии, что научный результат будет сформулирован как измеримый метод синхронизации и проверен сравнительным экспериментом.",
        default_bold=True,
        default_color=DARK,
    )

    for _ in range(2):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(8)

    metadata = (
        ("Проект", "«Контентфабрика» / OpenShorts"),
        ("Корпус", "16 русскоязычных + 19 англоязычных научных публикаций; 3 стандарта"),
        ("Назначение", "Основа первой главы ВКР и обзорной/методической статьи"),
        ("Дата фиксации", "23 июля 2026 г."),
        ("Версия", "1.0"),
    )
    for key, value in metadata:
        p = doc.add_paragraph(style="Report Meta")
        r = p.add_run(f"{key.upper():<18}")
        set_run_font(r, size=9.5, color=ACCENT)
        r.bold = True
        r2 = p.add_run(value)
        set_run_font(r2, size=10, color=TEXT)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(34)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Екатеринбург • 2026")
    set_run_font(r, size=10, color=MUTED)

    doc.add_page_break()
    toc_heading = doc.add_paragraph("Содержание", style="Heading 1")
    toc_heading.paragraph_format.space_before = Pt(0)
    contents = (
        "Резюме и решение о целесообразности темы",
        "1. Предметная область и границы исследования",
        "2. Методика поиска и отбора литературы",
        "3. Какие задачи решаются в предметной области",
        "4. Методы решения и их связь с OpenShorts",
        "5. Тематический синтез литературы",
        "6. Аннотированный список русскоязычных источников",
        "7. Аннотированный список англоязычных источников",
        "8. Исследовательский пробел",
        "9. Предлагаемый научный аппарат ВКР",
        "10. План экспериментального исследования",
        "11. Как превратить обзор в первую главу и статью",
        "12. Риски и угрозы валидности",
        "Заключение",
    )
    for item in contents:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.12)
        p.paragraph_format.right_indent = Inches(0.08)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_together = True
        add_inline(p, item, default_color=(ACCENT if item[0].isdigit() else DARK))
        set_paragraph_border(p, side="bottom", color="E6ECF2", size="3", space="3")
    doc.add_page_break()


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        values = [cell.strip() for cell in lines[i].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", value) for value in values):
            rows.append(values)
        i += 1
    return rows, i


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    set_table_width_and_indent(table)
    set_table_borders(table)

    if cols == 4:
        widths = [1500, 2800, 2100, 2960]
    elif cols == 3:
        widths = [2100, 3360, 3900]
    elif cols == 2:
        widths = [2400, 6960]
    else:
        widths = [9360 // cols] * cols

    for r_idx, row in enumerate(rows):
        row_obj = table.rows[r_idx]
        prevent_row_split(row_obj)
        if r_idx == 0:
            set_repeat_table_header(row_obj)
        for c_idx in range(cols):
            cell = row_obj.cells[c_idx]
            cell.width = Inches(widths[c_idx] / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_shading(cell, ACCENT if r_idx == 0 else (LIGHT if r_idx % 2 == 0 else WHITE))
            cell.text = ""
            p = cell.paragraphs[0]
            p.style = doc.styles["Table Text"]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            value = row[c_idx] if c_idx < len(row) else ""
            add_inline(p, value, default_bold=(r_idx == 0), default_color=(WHITE if r_idx == 0 else TEXT))
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(1)


def add_horizontal_rule(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(9)
    set_paragraph_border(p, side="bottom", color=BORDER, size="6", space="4")


def render_markdown(doc: Document, markdown: str) -> None:
    lines = markdown.splitlines()
    body_start = next((i + 1 for i, line in enumerate(lines) if line.strip() == "---"), 0)
    lines = lines[body_start:]
    i = 0
    paragraph_buffer: list[str] = []

    def flush_paragraph() -> None:
        nonlocal paragraph_buffer
        if paragraph_buffer:
            text = " ".join(part.strip() for part in paragraph_buffer)
            p = doc.add_paragraph()
            add_inline(p, text)
            paragraph_buffer = []

    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()

        if not stripped:
            flush_paragraph()
            i += 1
            continue

        if stripped.startswith("|"):
            flush_paragraph()
            rows, i = parse_table(lines, i)
            add_table(doc, rows)
            continue

        heading = re.match(r"^(#{1,3})\s+(.+)$", stripped)
        if heading:
            flush_paragraph()
            level = len(heading.group(1))
            text = heading.group(2)
            p = doc.add_paragraph(style=f"Heading {level}")
            add_inline(p, text)
            i += 1
            continue

        if stripped == "---":
            flush_paragraph()
            add_horizontal_rule(doc)
            i += 1
            continue

        if stripped.startswith("> "):
            flush_paragraph()
            quote_parts = [stripped[2:]]
            i += 1
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote_parts.append(lines[i].strip().lstrip(">").strip())
                i += 1
            p = doc.add_paragraph(style="Callout")
            shade_paragraph(p, PALE_BLUE)
            set_paragraph_border(p, side="left", color=ACCENT, size="18", space="8")
            add_inline(p, " ".join(quote_parts))
            continue

        bullet = re.match(r"^-\s+(.+)$", stripped)
        if bullet:
            flush_paragraph()
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, bullet.group(1))
            i += 1
            continue

        numbered = re.match(r"^\d+\.\s+(.+)$", stripped)
        if numbered:
            flush_paragraph()
            p = doc.add_paragraph(style="List Number")
            add_inline(p, numbered.group(1))
            i += 1
            continue

        paragraph_buffer.append(stripped)
        i += 1

    flush_paragraph()


def apply_document_defaults(doc: Document) -> None:
    settings = doc.settings._element
    compat = settings.find(qn("w:compat"))
    if compat is None:
        compat = OxmlElement("w:compat")
        settings.append(compat)
    setting = OxmlElement("w:compatSetting")
    setting.set(qn("w:name"), "compatibilityMode")
    setting.set(qn("w:uri"), "http://schemas.microsoft.com/office/word")
    setting.set(qn("w:val"), "15")
    compat.append(setting)

    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "false")


def build() -> None:
    markdown = SOURCE.read_text(encoding="utf-8")
    doc = Document()
    create_styles(doc)
    configure_sections(doc)
    apply_document_defaults(doc)
    add_cover(doc)
    render_markdown(doc, markdown)

    props = doc.core_properties
    props.title = "Анализ литературы: мультимодальная генерация коротких вертикальных видеороликов"
    props.subject = "Рабочая основа первой главы ВКР и научной статьи"
    props.author = "OpenShorts"
    props.keywords = "ВКР, генеративное видео, AI Shorts, ASR, TTS, lip-sync, мультимодальная синхронизация"
    props.comments = "Сформировано из верифицированного корпуса источников по состоянию на 23.07.2026."

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
